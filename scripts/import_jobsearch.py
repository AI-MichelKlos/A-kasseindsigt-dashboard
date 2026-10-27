"""Læs manuelle Jobindsats-eksporter af dagpengemodtageres jobsøgning.

Rapporten har ikke et API-tabel. Filerne ligger i data/raw/jobsearch/ som
YYYYMM-fordeling.xlsx og YYYYMM-tabel.xlsx. En måned uden begge filer
springes over, så den øvrige opdatering kan fortsætte.
"""

from __future__ import annotations

import json
import math
import re
import sys
import warnings
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import jobindsats_api as ji

warnings.filterwarnings("ignore", message="Workbook contains no default style")

BASE = Path(__file__).resolve().parents[1]
DATA = BASE / "data" / "dashboard-data.json"
NAMES = BASE / "config" / "a-kasse-navne.json"
RAW = BASE / "data" / "raw" / "jobsearch"
REPORT_URL = "https://jobindsats.dk/rapportbank/rapporter/dagpengemodtageres-jobsogning/"
DATASET = "rapport:dagpengemodtageres-jobsogning"
STATUS_KEY = "jobJobSearch"
NOTE = (
    "Jobsøgning blandt dagpengemodtagere hentes fra manuelle eksporter af Jobindsats-rapporten "
    "Dagpengemodtageres jobsøgning og indgår ikke i den automatiske API-opdatering."
)
SERIES_KEYS = (
    "shareWithApplication",
    "shareWithoutApplication",
    "countWithApplication",
    "countWithoutApplication",
    "basePopulation",
    "avgApplications",
)
MONTHS = {
    "januar": 1,
    "februar": 2,
    "marts": 3,
    "april": 4,
    "maj": 5,
    "juni": 6,
    "juli": 7,
    "august": 8,
    "september": 9,
    "oktober": 10,
    "november": 11,
    "december": 12,
}
PERIOD_RE = re.compile(r"Periode er ([A-Za-zÆØÅæøå]+) (\d{4})")
FILE_RE = re.compile(r"^(\d{6})-(fordeling|tabel)\.xlsx$")
SHARE_TOLERANCE = 1e-6


def now_iso():
    return datetime.now(ZoneInfo("Europe/Copenhagen")).isoformat(timespec="seconds")


def load_workbook_rows(path: Path):
    import openpyxl

    # read_only stopper ved den første celle i disse Power BI-eksporter.
    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        sheet = workbook.active
        return [tuple(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def cell(row, index):
    if index is None or index >= len(row):
        return None
    return row[index]


def header_index(rows, wanted, limit=8):
    target = ji.norm(wanted)
    for index, row in enumerate(rows[:limit]):
        if row and ji.norm(row[0]) == target:
            return index
    return None


def column_index(header, *names):
    targets = {ji.norm(name) for name in names}
    for index, value in enumerate(header):
        if ji.norm(value) in targets:
            return index
    return None


def period_from_text(text):
    match = PERIOD_RE.search(str(text or ""))
    if not match:
        return None
    month = MONTHS.get(match.group(1).lower())
    if not month:
        raise RuntimeError(f"Ukendt månedsnavn i periodefeltet: {match.group(1)}")
    return f"{int(match.group(2)):04d}M{month:02d}"


def filename_period(stamp):
    year, month = int(stamp[:4]), int(stamp[4:])
    if month < 1 or month > 12:
        raise RuntimeError(f"Ugyldigt månedsnummer i filnavn: {stamp}")
    return f"{year:04d}M{month:02d}"


def require_period(text, expected, path: Path):
    found = period_from_text(text)
    if found and found != expected:
        raise RuntimeError(f"{path.name} angiver {found}, men filnavnet er {expected}")
    return found


def require_area(text, path: Path):
    if "Område er Alle kommuner" not in str(text or ""):
        raise RuntimeError(f"{path.name} er ikke afgrænset til Område = Alle kommuner")


def plain_number(value):
    """Læs et tal uden at afrunde brøken, før den omregnes til procent."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, str):
        text = value.strip().replace("\u00a0", "").replace(" ", "")
        if text.lower() in {"", "-", ".", "..", "null", "none", "nan"}:
            return None
        if "," in text:
            text = text.replace(".", "").replace(",", ".")
        try:
            value = float(text)
        except ValueError:
            return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(number):
        return None
    return number


def percent(value):
    number = plain_number(value)
    if number is None:
        return None
    return round(number * 100, 6)


def count(value):
    number = plain_number(value)
    if number is None:
        return None
    if number.is_integer():
        return int(number)
    return round(number, 6)


def build_aliases(data):
    total = data["meta"]["totalFundCode"]
    owners = defaultdict(dict)

    def bind(label, code, origin):
        key = ji.norm(label)
        if not key or not code:
            return
        owners[key].setdefault(code, origin)

    for label in ("Alle A-kasser", "I alt", "A-kasser i alt"):
        bind(label, total, "total")

    reference = json.loads(NAMES.read_text(encoding="utf-8"))
    by_short = {}
    by_name = {}
    for code, fund in data.get("funds", {}).items():
        if code == total:
            continue
        short = ji.norm(fund.get("short"))
        name = ji.norm(fund.get("name"))
        if short:
            by_short.setdefault(short, code)
        if name:
            by_name.setdefault(name, code)

    for row in reference.get("funds", []):
        code = by_short.get(ji.norm(row.get("dakShort"))) or by_name.get(ji.norm(row.get("dakName")))
        if not code:
            continue
        for field in ("dakName", "dakShort", "dstName", "jobindsatsName", "starMemberName"):
            bind(row.get(field), code, field)

    for code, fund in data.get("funds", {}).items():
        if code == total:
            continue
        bind(fund.get("short"), code, "fund.short")
        bind(fund.get("name"), code, "fund.name")
        for key, value in (fund.get("sourceNames") or {}).items():
            bind(value, code, f"sourceNames.{key}")

    aliases = {}
    ambiguous = []
    for key, codes in owners.items():
        if len(codes) == 1:
            aliases[key] = next(iter(codes))
        else:
            ambiguous.append({"label": key, "codes": sorted(codes)})
    return aliases, ambiguous


def resolve_label(label, aliases, bridge):
    key = ji.norm(label)
    if not key:
        return None, "tomt navn"
    direct = aliases.get(key)
    linked = bridge.get(key)
    if direct and linked and direct != linked:
        return None, f"flertydigt navn {label}"
    code = direct or linked
    if not code:
        return None, None
    return code, None


def read_fordeling(path: Path, expected, aliases):
    rows = load_workbook_rows(path)
    header_at = header_index(rows, "A-kasse forkortelse")
    if header_at is None:
        raise RuntimeError(f"{path.name} mangler kolonnen A-kasse forkortelse")
    require_area(rows[0][0] if rows else "", path)
    require_period(rows[0][0] if rows else "", expected, path)
    header = rows[header_at]
    columns = {
        "abbr": column_index(header, "A-kasse forkortelse"),
        "name": column_index(header, "Første A-kasse"),
        "population": column_index(header, "Grundpopulation"),
        "average_share": column_index(header, "Gennemsnit alle a-kasser"),
        "share": column_index(header, "Andel"),
    }
    missing = [key for key, index in columns.items() if index is None]
    if missing:
        raise RuntimeError(f"{path.name} mangler kolonner: {missing}")
    avg_index = next(
        (
            index
            for index, value in enumerate(header)
            if ji.norm(value).startswith("gennemsnitlige antal jobsoegning")
        ),
        None,
    )
    parsed = {}
    bridge = {}
    unmapped = []
    average_share = None
    for row in rows[header_at + 1 :]:
        abbr = cell(row, columns["abbr"])
        name = cell(row, columns["name"])
        if abbr in (None, "") and name in (None, ""):
            continue
        if str(abbr or "").startswith("Anvendte filtre"):
            break
        code_by_abbr = aliases.get(ji.norm(abbr))
        code_by_name = aliases.get(ji.norm(name))
        if code_by_abbr and code_by_name and code_by_abbr != code_by_name:
            unmapped.append(f"{abbr} / {name} (forkortelse og navn peger på forskellige a-kasser)")
            continue
        code = code_by_abbr or code_by_name
        if not code:
            unmapped.append(str(abbr or name))
            continue
        share = percent(cell(row, columns["share"]))
        population = count(cell(row, columns["population"]))
        applications = count(cell(row, avg_index)) if avg_index is not None else None
        observed_average = percent(cell(row, columns["average_share"]))
        if observed_average is None:
            raise RuntimeError(f"{path.name} mangler gennemsnit for alle a-kasser på rækken {abbr or name}")
        if average_share is None:
            average_share = observed_average
        elif abs(average_share - observed_average) > SHARE_TOLERANCE:
            raise RuntimeError(f"{path.name} har forskellige værdier i Gennemsnit alle a-kasser")
        parsed[code] = {
            "shareWithApplication": share,
            "basePopulation": population,
            "avgApplications": applications,
            "sourceShare": plain_number(cell(row, columns["share"])),
        }
        for label in (abbr, name):
            key = ji.norm(label)
            if not key:
                continue
            previous = bridge.get(key)
            if previous and previous != code:
                unmapped.append(f"{label} (brugt om både {previous} og {code} i {path.name})")
            else:
                bridge[key] = code
    if average_share is None:
        raise RuntimeError(f"{path.name} har ingen a-kasserækker")
    return parsed, bridge, unmapped, average_share, avg_index is not None


def read_tabel(path: Path, expected, aliases, bridge):
    rows = load_workbook_rows(path)
    header_at = header_index(rows, "A-kasse")
    if header_at is None or header_at == 0:
        raise RuntimeError(f"{path.name} mangler kolonneoverskriften A-kasse")
    category = rows[header_at - 1]
    if "uden" not in ji.norm(cell(category, 1)) or "mindst" not in ji.norm(cell(category, 3)):
        raise RuntimeError(f"{path.name} har et uventet kategoriheader: {category}")
    footer = next((row[0] for row in rows if row and str(row[0] or "").startswith("Anvendte filtre")), "")
    if footer:
        require_area(footer, path)
        require_period(footer, expected, path)
    parsed = {}
    unmapped = []
    for row in rows[header_at + 1 :]:
        label = cell(row, 0)
        if label in (None, ""):
            continue
        text = str(label)
        if text.startswith("Anvendte filtre"):
            break
        code, reason = resolve_label(text, aliases, bridge)
        if reason:
            unmapped.append(f"{text} ({reason})")
            continue
        if not code:
            unmapped.append(text)
            continue
        parsed[code] = {
            "countWithoutApplication": count(cell(row, 1)),
            "shareWithoutApplication": percent(cell(row, 2)),
            "countWithApplication": count(cell(row, 3)),
            "shareWithApplication": percent(cell(row, 4)),
            "sourceShare": plain_number(cell(row, 4)),
        }
    return parsed, unmapped


def check_shares(period, code, left, right, label):
    if left is None or right is None:
        return
    if abs(left - right) > SHARE_TOLERANCE:
        raise RuntimeError(
            f"{period} {label}: Andel i fordeling ({left}) stemmer ikke med tabellen ({right}) for {code}"
        )


def check_population(period, code, population, with_count, without_count):
    if None in (population, with_count, without_count):
        return
    if int(population) != int(with_count) + int(without_count):
        raise RuntimeError(
            f"{period} {code}: grundpopulation {population} er ikke lig med antal med ({with_count}) "
            f"plus antal uden ({without_count})"
        )


def discover_files():
    grouped = defaultdict(dict)
    ignored = []
    if not RAW.is_dir():
        return {}, ignored
    for path in sorted(RAW.glob("*.xlsx")):
        match = FILE_RE.fullmatch(path.name)
        if not match:
            ignored.append(path.name)
            continue
        grouped[match.group(1)][match.group(2)] = path
    return grouped, ignored


def write_data(data):
    DATA.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def mark_missing(data, message):
    for fund in data.get("funds", {}).values():
        fund.get("jobindsats", {}).pop("jobSearch", None)
    data.setdefault("meta", {}).setdefault("sourceStatus", {})[STATUS_KEY] = {
        "state": "not_configured",
        "source": "Jobindsats.dk / STAR",
        "dataset": DATASET,
        "unit": "pct. dagpengemodtagere med mindst én jobansøgning i måneden",
        "error": message,
        "checkedAt": now_iso(),
        "reportUrl": REPORT_URL,
    }
    notes = data["meta"].setdefault("methodNotes", [])
    if NOTE not in notes:
        notes.append(NOTE)
    write_data(data)


def main():
    data = json.loads(DATA.read_text(encoding="utf-8"))
    grouped, ignored = discover_files()
    complete = {stamp: files for stamp, files in grouped.items() if {"fordeling", "tabel"} <= set(files)}
    incomplete = sorted(stamp for stamp in grouped if stamp not in complete)
    if not complete:
        message = "Ingen komplette månedsfiler i data/raw/jobsearch."
        if incomplete:
            message += " Ufuldstændige måneder: " + ", ".join(incomplete) + "."
        mark_missing(data, message)
        print(message)
        return 0

    aliases, ambiguous = build_aliases(data)
    total_code = data["meta"]["totalFundCode"]
    by_period = {}
    unmapped = []
    has_average_applications = False
    for stamp in sorted(complete):
        period = filename_period(stamp)
        fordeling, bridge, fordeling_unmapped, average_share, has_avg = read_fordeling(
            complete[stamp]["fordeling"], period, aliases
        )
        tabel, tabel_unmapped = read_tabel(complete[stamp]["tabel"], period, aliases, bridge)
        has_average_applications = has_average_applications or has_avg
        for label in fordeling_unmapped + tabel_unmapped:
            unmapped.append(f"{period}: {label}")
        if total_code not in tabel:
            raise RuntimeError(f"{complete[stamp]['tabel'].name} mangler rækken Alle A-kasser")
        total_share = tabel[total_code]["sourceShare"]
        check_shares(period, total_code, average_share / 100, total_share, "Gennemsnit alle a-kasser")
        codes = set(fordeling) | set(tabel)
        merged = {}
        for code in codes:
            left = fordeling.get(code, {})
            right = tabel.get(code, {})
            check_shares(period, code, left.get("sourceShare"), right.get("sourceShare"), "Andel")
            share = right.get("shareWithApplication", left.get("shareWithApplication"))
            record = {
                "shareWithApplication": share,
                "shareWithoutApplication": right.get("shareWithoutApplication"),
                "countWithApplication": right.get("countWithApplication"),
                "countWithoutApplication": right.get("countWithoutApplication"),
                "basePopulation": left.get("basePopulation"),
                "avgApplications": left.get("avgApplications"),
            }
            check_population(
                period,
                code,
                record["basePopulation"],
                record["countWithApplication"],
                record["countWithoutApplication"],
            )
            merged[code] = record
        by_period[period] = merged

    periods = sorted(by_period, key=lambda value: (int(value[:4]), int(value[5:])))
    present_codes = {code for period in by_period.values() for code in period}
    for code, fund in data.get("funds", {}).items():
        block = fund.setdefault("jobindsats", {}).get("jobSearch")
        if code not in present_codes:
            if isinstance(block, dict):
                fund["jobindsats"].pop("jobSearch", None)
            continue
        series = {key: [] for key in SERIES_KEYS}
        for period in periods:
            record = by_period[period].get(code, {})
            for key in SERIES_KEYS:
                series[key].append(record.get(key))
        fund["jobindsats"]["jobSearch"] = {"labels": periods, **series}

    latest_file = sorted(complete)[-1]
    data["meta"].setdefault("sourceStatus", {})[STATUS_KEY] = {
        "state": "ok",
        "source": "Jobindsats.dk / STAR",
        "dataset": DATASET,
        "latestPeriod": periods[-1],
        "unit": "pct. dagpengemodtagere med mindst én jobansøgning i måneden",
        "note": (
            "Manuel eksport i data/raw/jobsearch fra rapporten Dagpengemodtageres jobsøgning. "
            f"Der er ikke et Jobindsats API-tabel for målingen. Seneste filer: {latest_file}-fordeling.xlsx "
            f"og {latest_file}-tabel.xlsx."
            + ("" if has_average_applications else " Eksporten havde ikke kolonnen med gennemsnitligt antal jobsøgninger.")
        ),
        "checkedAt": now_iso(),
        "reportUrl": REPORT_URL,
        "manual": True,
    }
    notes = data["meta"].setdefault("methodNotes", [])
    if NOTE not in notes:
        notes.append(NOTE)
    write_data(data)

    funds = data["funds"]
    mapped = []
    for code in sorted(present_codes, key=lambda item: (item != total_code, ji.norm(funds.get(item, {}).get("name", item)))):
        fund = funds.get(code, {})
        mapped.append(
            {
                "code": code,
                "short": fund.get("short") or fund.get("name") or code,
                "name": fund.get("name") or code,
            }
        )
    missing_funds = [
        {"code": code, "name": fund.get("name", code)}
        for code, fund in funds.items()
        if code not in present_codes
    ]
    total = by_period[periods[-1]][total_code]
    summary = {
        "latestPeriod": periods[-1],
        "months": periods,
        "incompleteMonths": [filename_period(stamp) for stamp in incomplete],
        "ignoredFiles": ignored,
        "ambiguousAliases": ambiguous,
        "mapped": mapped,
        "unmapped": unmapped,
        "fundsWithoutSeries": missing_funds,
        "latestTotal": {
            "countWithApplication": total.get("countWithApplication"),
            "shareWithApplication": total.get("shareWithApplication"),
            "countWithoutApplication": total.get("countWithoutApplication"),
            "shareWithoutApplication": total.get("shareWithoutApplication"),
        },
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:
        print(f"Jobsøgning blev ikke importeret: {exc}", file=sys.stderr)
        sys.exit(1)
